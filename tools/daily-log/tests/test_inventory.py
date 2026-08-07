import datetime as dt
import unittest

from daily_log.inventory import (
    Candidate,
    build_candidates,
    evidence_note,
    normalize,
    suggest_frequency,
)


def record(subject, day, *, content="", kind="메일수신", counterpart="a@x.com"):
    return {
        "at": dt.datetime(2026, 8, day, 9, 0).isoformat(),
        "date": dt.date(2026, 8, day).isoformat(),
        "kind": kind,
        "counterpart": counterpart,
        "subject": subject,
        "content": content or subject,
    }


class 제목_정규화(unittest.TestCase):
    def test_회신_전달_표식을_뗀다(self):
        self.assertEqual(normalize("RE: 정산 자료"), normalize("정산 자료"))
        self.assertEqual(normalize("FW: 정산 자료"), normalize("정산 자료"))
        self.assertEqual(normalize("RE: RE: FW: 정산 자료"), normalize("정산 자료"))

    def test_괄호_안_고유값을_뗀다(self):
        # 이게 핵심이다. 매번 이름과 날짜가 바뀌어도 같은 업무로 묶여야 한다.
        first = normalize("VPN 접속 승인요청서(홍길동,26.08.06)")
        second = normalize("VPN 접속 승인요청서(임꺽정,26.09.02)")
        self.assertEqual(first, second)

    def test_날짜가_붙어도_묶인다(self):
        self.assertEqual(normalize("MPS-MRP 수행의 건 26.08.06"), normalize("MPS-MRP 수행의 건"))

    def test_다른_업무는_안_묶인다(self):
        self.assertNotEqual(normalize("VPN 접속 승인요청서"), normalize("전기 사용량 공지"))

    def test_빈_제목은_빈_키다(self):
        self.assertEqual(normalize(""), "")
        self.assertEqual(normalize("   "), "")


class 후보_묶기(unittest.TestCase):
    def test_보낸_것과_회신을_한_업무로_묶는다(self):
        records = [
            record("(중요) ISIR 공정감사 체크리스트 송부의 건", 6, kind="메일발신"),
            record("RE: (중요) ISIR 공정감사 체크리스트 송부의 건", 6),
        ]
        candidates = build_candidates(records)
        self.assertEqual(len(candidates), 1)
        self.assertEqual(candidates[0].count, 2)

    def test_많이_나온_업무가_먼저_온다(self):
        records = [record("드문 일", 6)] + [record("잦은 일", d) for d in (6, 7, 8)]
        self.assertEqual([c.name for c in build_candidates(records)], ["잦은 일", "드문 일"])

    def test_메모는_작업_절차로_모인다(self):
        records = [
            record("VPN 접속 승인요청서(홍길동)", 6, content="검토 후 도장 날인"),
            record("VPN 접속 승인요청서(임꺽정)", 20, content="상부 전달"),
        ]
        candidate = build_candidates(records)[0]
        self.assertIn("검토 후 도장 날인", candidate.procedure)
        self.assertIn("상부 전달", candidate.procedure)

    def test_같은_메모가_반복돼도_한_번만_적는다(self):
        records = [record("승인 요청", d, content="검토 후 날인") for d in (6, 13, 20)]
        self.assertEqual(build_candidates(records)[0].procedure, "검토 후 날인")

    def test_제목이_없는_옛_기록은_content_로_묶는다(self):
        old = {"at": dt.datetime(2026, 8, 6, 9).isoformat(), "kind": "메일수신",
               "counterpart": "a@x.com", "content": "예전 기록"}
        self.assertEqual(build_candidates([old])[0].name, "예전 기록")

    def test_내용이_비면_버린다(self):
        self.assertEqual(build_candidates([record("", 6, content="")]), [])

    def test_업무명에서_그_건의_고유값을_뗀다(self):
        # 인벤토리에 들어갈 이름이다. 특정인 이름과 날짜가 남으면 반복업무 이름이 아니다.
        records = [record("FW: VPN 접속 승인요청서(홍길동,26.08.06)", 6)]
        self.assertEqual(build_candidates(records)[0].name, "VPN 접속 승인요청서")

    def test_식별번호_말머리는_떼고_뜻있는_말머리는_남긴다(self):
        cases = {
            "RE:[CASE 000000000000000] SES: Production Access": "SES: Production Access",
            "[테크톡] 보안 인증 경험": "[테크톡] 보안 인증 경험",
        }
        for subject, expected in cases.items():
            self.assertEqual(build_candidates([record(subject, 6)])[0].name, expected)

    def test_다_떼면_남는게_없을_때는_원문을_쓴다(self):
        self.assertEqual(build_candidates([record("[CASE 000000000000000]", 6)])[0].name,
                         "[CASE 000000000000000]")


class 주기_추정(unittest.TestCase):
    def _candidate(self, days):
        return Candidate(
            name="x",
            occurrences=[dt.datetime(2026, 8, 1) + dt.timedelta(days=d) for d in days],
        )

    def test_관측이_짧으면_주기를_말하지_않는다(self):
        # 이게 이 모듈의 약속이다. 짧게 보고 단정하면 우선순위 전체가 틀어진다.
        self.assertEqual(suggest_frequency(self._candidate([0, 1, 2, 3, 4])), "")

    def test_한_번만_나온_것은_주기가_없다(self):
        self.assertEqual(suggest_frequency(self._candidate([0])), "")

    def test_매일_나오면_매일로_본다(self):
        self.assertEqual(suggest_frequency(self._candidate(list(range(0, 28)))), "매일")

    def test_주_1회면_주_1회로_본다(self):
        self.assertEqual(suggest_frequency(self._candidate([0, 7, 14, 21, 27])), "주 1회")

    def test_월_1회는_한_달을_봐야_말한다(self):
        month = self._candidate([0, 30, 60])
        self.assertEqual(suggest_frequency(month), "월 1회")

    def test_근거가_없으면_메모가_이유를_밝힌다(self):
        note = evidence_note(self._candidate([0, 1]))
        self.assertIn("주기는 관측", note)


class 분류_추정(unittest.TestCase):
    def test_승인_업무는_문서_승인으로_본다(self):
        c = Candidate(name="VPN 접속 승인요청서", memos=["검토 후 도장 날인"])
        self.assertEqual(c.category, "문서·승인")

    def test_장애_대응은_시스템_운영으로_본다(self):
        self.assertEqual(Candidate(name="PC점검 아웃룩 색인 오류").category, "시스템 운영")

    def test_모르면_기타로_둔다(self):
        self.assertEqual(Candidate(name="ㅁㄴㅇㄹ").category, "기타")

    def test_시스템_이름을_건진다(self):
        c = Candidate(name="PC점검", memos=["아웃룩 재색인 진행"])
        self.assertIn("아웃룩", c.systems)


class 인벤토리에_쓰기(unittest.TestCase):
    def setUp(self):
        import pathlib
        import tempfile

        from openpyxl import Workbook

        from daily_log import inventory

        wb = Workbook()
        wb.active.title = inventory.SHEET
        self.path = pathlib.Path(tempfile.mkdtemp()) / "inv.xlsx"
        wb.save(self.path)

    def _read(self, row, column):
        from openpyxl import load_workbook

        from daily_log import inventory

        return load_workbook(self.path)[inventory.SHEET].cell(row=row, column=column).value

    def _write(self, candidates):
        from daily_log import inventory

        return inventory.write_into_inventory(self.path, candidates)

    def test_후보를_첫_빈_행부터_채운다(self):
        from daily_log import inventory

        added, _, _ = self._write(build_candidates([record("승인 요청", 6)]))
        self.assertEqual(added, ["승인 요청"])
        self.assertEqual(self._read(inventory.ROW_FIRST, inventory.COL_NAME), "승인 요청")

    def test_소요시간은_비워둔다(self):
        # 이 모듈의 약속. 여기 숫자가 들어가면 우선순위 전체가 내 감이 된다.
        from daily_log import inventory

        self._write(build_candidates([record("승인 요청", 6)]))
        self.assertIsNone(self._read(inventory.ROW_FIRST, inventory.COL_MINUTES))

    def test_이미_있는_업무는_건드리지_않는다(self):
        from openpyxl import load_workbook

        from daily_log import inventory

        wb = load_workbook(self.path)
        ws = wb[inventory.SHEET]
        ws.cell(row=inventory.ROW_FIRST, column=inventory.COL_NAME, value="승인 요청")
        ws.cell(row=inventory.ROW_FIRST, column=inventory.COL_MINUTES, value=45)
        wb.save(self.path)

        added, skipped, _ = self._write(build_candidates([record("RE: 승인 요청", 6)]))
        self.assertEqual(added, [])
        self.assertEqual(skipped, ["승인 요청"])
        # 사용자가 채운 소요시간이 살아 있어야 한다
        self.assertEqual(self._read(inventory.ROW_FIRST, inventory.COL_MINUTES), 45)

    def test_업무명만_지운_행을_재사용해도_옛_값이_안_남는다(self):
        # openpyxl 은 cell(value=None) 을 무시한다. 그걸 모르고 쓰면 지운 업무의
        # 주기·절차가 새 업무에 그대로 붙는다.
        from openpyxl import load_workbook

        from daily_log import inventory

        wb = load_workbook(self.path)
        ws = wb[inventory.SHEET]
        ws.cell(row=inventory.ROW_FIRST, column=inventory.COL_FREQUENCY, value="매일")
        ws.cell(row=inventory.ROW_FIRST, column=inventory.COL_MINUTES, value=90)
        ws.cell(row=inventory.ROW_FIRST, column=inventory.COL_PROCEDURE, value="옛 절차")
        wb.save(self.path)

        self._write(build_candidates([record("새 업무", 6)]))
        self.assertEqual(self._read(inventory.ROW_FIRST, inventory.COL_NAME), "새 업무")
        for column in (inventory.COL_FREQUENCY, inventory.COL_MINUTES, inventory.COL_PROCEDURE):
            self.assertIsNone(self._read(inventory.ROW_FIRST, column))

    def test_같은_명령을_두_번_돌려도_행이_안_늘어난다(self):
        candidates = build_candidates([record("승인 요청", 6)])
        self.assertEqual(self._write(candidates)[0], ["승인 요청"])
        self.assertEqual(self._write(candidates)[0], [])

    def test_시트가_없으면_분명히_실패한다(self):
        import pathlib
        import tempfile

        from openpyxl import Workbook

        from daily_log import inventory

        wb = Workbook()
        wb.active.title = "엉뚱한시트"
        path = pathlib.Path(tempfile.mkdtemp()) / "wrong.xlsx"
        wb.save(path)
        with self.assertRaises(ValueError):
            inventory.write_into_inventory(path, build_candidates([record("x", 6)]))


if __name__ == "__main__":
    unittest.main()
