import datetime as dt
import pathlib
import tempfile
import unittest

from openpyxl import load_workbook

from daily_log import queue_sheet, weekly
from daily_log.models import MAIL_RECEIVED, QueueItem
from daily_log.rules import LearnedStats
from daily_log.store import Store

DAY = dt.date(2026, 8, 6)


def item(key, subject, at_hour=9, counterpart="a@x.com"):
    return QueueItem(
        key=key,
        at=dt.datetime.combine(DAY, dt.time(at_hour)),
        kind=MAIL_RECEIVED,
        counterpart=counterpart,
        subject=subject,
        reason="내가 수신(To)",
    )


class 대기열_엑셀(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = pathlib.Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_쓰고_다시_읽으면_그대로다(self):
        path = queue_sheet.write_queue(
            self.dir / "q.xlsx", DAY, [item("k1", "첫째"), item("k2", "둘째", 10)]
        )
        self.assertEqual(queue_sheet.read_kept(path), {"k1": "첫째", "k2": "둘째"})

    def test_행을_지우면_읽히지_않는다(self):
        # 이게 이 도구의 핵심 상호작용이다. 깨지면 전부 무의미해진다.
        path = queue_sheet.write_queue(
            self.dir / "q.xlsx", DAY, [item("k1", "남길 것"), item("k2", "지울 것", 10)]
        )
        wb = load_workbook(path)
        ws = wb["오늘 업무"]
        for row in range(ws.max_row, queue_sheet.FIRST_ROW - 1, -1):
            if ws.cell(row=row, column=4).value == "지울 것":
                ws.delete_rows(row)
        wb.save(path)

        kept = queue_sheet.read_kept(path)
        self.assertIn("k1", kept)
        self.assertNotIn("k2", kept)

    def test_업무내용을_적으면_제목_대신_그게_들어간다(self):
        path = queue_sheet.write_queue(self.dir / "q.xlsx", DAY, [item("k1", "RE: RE: 회신")])
        wb = load_workbook(path)
        wb["오늘 업무"].cell(
            row=queue_sheet.FIRST_ROW, column=queue_sheet.CONTENT_COLUMN, value="정산 자료 회신"
        )
        wb.save(path)
        self.assertEqual(queue_sheet.read_kept(path)["k1"], "정산 자료 회신")

    def test_구분을_아님으로_바꾸면_읽히지_않는다(self):
        # 행을 지우는 대신 표시만 하는 방식. 왜 아닌지가 메모로 남는다.
        path = queue_sheet.write_queue(
            self.dir / "q.xlsx", DAY, [item("k1", "업무"), item("k2", "광고", 10)]
        )
        wb = load_workbook(path)
        ws = wb["오늘 업무"]
        for row in range(queue_sheet.FIRST_ROW, ws.max_row + 1):
            if ws.cell(row=row, column=4).value == "광고":
                ws.cell(row=row, column=queue_sheet.MARK_COLUMN, value="아님")
                ws.cell(row=row, column=queue_sheet.CONTENT_COLUMN, value="광고성 메일")
        wb.save(path)

        kept = queue_sheet.read_kept(path)
        self.assertIn("k1", kept)
        self.assertNotIn("k2", kept)

    def test_아님_표기가_달라도_받아준다(self):
        # 사람마다 X, 제외, 아님을 섞어 쓴다. 표기 하나 때문에 업무가 새면 안 된다.
        for mark in ("아님", "X", "x", "제외"):
            path = queue_sheet.write_queue(self.dir / f"q{mark}.xlsx", DAY, [item("k1", "광고")])
            wb = load_workbook(path)
            wb["오늘 업무"].cell(
                row=queue_sheet.FIRST_ROW, column=queue_sheet.MARK_COLUMN, value=mark
            )
            wb.save(path)
            self.assertNotIn("k1", queue_sheet.read_kept(path), f"표기 {mark!r} 를 놓쳤다")

    def test_구분이_비어있으면_업무로_본다(self):
        # 지우는 방식만 쓰던 사람이 구분 칸을 몰라도 그대로 동작해야 한다.
        path = queue_sheet.write_queue(self.dir / "q.xlsx", DAY, [item("k1", "본문")])
        self.assertIn("k1", queue_sheet.read_kept(path))

    def test_자동제외_시트에_그대로_두면_계속_빠진다(self):
        # 예전엔 두 시트를 다 읽어서, 가만히 두면 자동 제외가 하루 만에 풀렸다.
        path = queue_sheet.write_queue(
            self.dir / "q.xlsx", DAY, [item("k1", "본문")], [item("k2", "자동제외됨", 11)]
        )
        self.assertEqual(set(queue_sheet.read_kept(path)), {"k1"})

    def test_자동제외_시트에서_옮겨오면_되살아난다(self):
        # 잘못 제외된 걸 사용자가 되살릴 수 있어야 한다 — 안내문이 시키는 대로 옮겨본다.
        path = queue_sheet.write_queue(
            self.dir / "q.xlsx", DAY, [item("k1", "본문")], [item("k2", "자동제외됨", 11)]
        )
        wb = load_workbook(path)
        source, target = wb["자동 제외됨"], wb["오늘 업무"]
        moved = [c.value for c in source[queue_sheet.FIRST_ROW]]
        for index, value in enumerate(moved, start=1):
            target.cell(row=target.max_row + (1 if index == 1 else 0), column=index, value=value)
        wb.save(path)

        self.assertEqual(set(queue_sheet.read_kept(path)), {"k1", "k2"})

    def test_아님으로_표시한_행의_메모도_남는다(self):
        # 왜 업무가 아닌지가 반복업무를 추릴 때 근거가 된다. 버리면 안 된다.
        path = queue_sheet.write_queue(self.dir / "q.xlsx", DAY, [item("k1", "광고")])
        wb = load_workbook(path)
        ws = wb["오늘 업무"]
        ws.cell(row=queue_sheet.FIRST_ROW, column=queue_sheet.MARK_COLUMN, value="아님")
        ws.cell(row=queue_sheet.FIRST_ROW, column=queue_sheet.CONTENT_COLUMN, value="광고성 메일")
        wb.save(path)

        self.assertEqual(queue_sheet.read_memos(path)["k1"], (False, "광고성 메일"))

    def test_자동제외가_없으면_시트를_만들지_않는다(self):
        path = queue_sheet.write_queue(self.dir / "q.xlsx", DAY, [item("k1", "본문")])
        self.assertEqual(load_workbook(path).sheetnames, ["오늘 업무"])


class 저장소(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.store = Store(pathlib.Path(self._tmp.name))

    def tearDown(self):
        self._tmp.cleanup()

    def test_스냅샷_왕복(self):
        self.store.save_snapshot(DAY, [item("k1", "제목")])
        restored = self.store.load_snapshot(DAY)
        self.assertEqual(len(restored), 1)
        self.assertEqual(restored[0].key, "k1")
        self.assertEqual(restored[0].at.hour, 9)

    def test_스냅샷이_없으면_빈_목록(self):
        self.assertEqual(self.store.load_snapshot(DAY), [])

    def test_같은_날을_다시_확정하면_덮어쓴다(self):
        # 하루에 두 번 돌려도 기록이 두 배가 되면 안 된다.
        self.store.append_daily(DAY, [{"at": "09:00", "content": "첫 번째"}])
        self.store.append_daily(DAY, [{"at": "09:00", "content": "다시 확정"}])
        records = self.store.load_daily()
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["content"], "다시 확정")

    def test_다른_날은_보존한다(self):
        other = DAY - dt.timedelta(days=1)
        self.store.append_daily(other, [{"at": "09:00", "content": "어제"}])
        self.store.append_daily(DAY, [{"at": "09:00", "content": "오늘"}])
        self.assertEqual(len(self.store.load_daily()), 2)

    def test_기간으로_거른다(self):
        self.store.append_daily(DAY - dt.timedelta(days=10), [{"content": "옛날"}])
        self.store.append_daily(DAY, [{"content": "이번주"}])
        records = self.store.load_daily(DAY - dt.timedelta(days=2), DAY)
        self.assertEqual([r["content"] for r in records], ["이번주"])

    def test_깨진_줄이_있어도_나머지는_살린다(self):
        self.store.append_daily(DAY, [{"content": "정상"}])
        with self.store.daily_path.open("a", encoding="utf-8") as f:
            f.write("{망가진 줄\n")
        self.assertEqual(len(self.store.load_daily()), 1)

    def test_학습_통계_왕복(self):
        stats = LearnedStats()
        stats.observe("spam@x.com", was_dropped=True)
        stats.observe("spam@x.com", was_dropped=True)
        self.store.save_stats(stats)
        self.assertEqual(self.store.load_stats().drop_rate("spam@x.com"), 1.0)

    def test_통계_파일이_깨져도_죽지_않는다(self):
        self.store.dir.mkdir(parents=True, exist_ok=True)
        self.store.stats_path.write_text("{깨짐", encoding="utf-8")
        self.assertEqual(self.store.load_stats().seen, {})


class 주간보고(unittest.TestCase):
    def test_주의_월요일과_금요일을_찾는다(self):
        start, end = weekly.week_bounds(dt.date(2026, 8, 6))  # 목요일
        self.assertEqual(start, dt.date(2026, 8, 3))
        self.assertEqual(end, dt.date(2026, 8, 7))

    def test_월요일에_실행해도_그_주가_나온다(self):
        start, _ = weekly.week_bounds(dt.date(2026, 8, 3))
        self.assertEqual(start, dt.date(2026, 8, 3))

    def test_기록이_있는_날만_내용이_들어간다(self):
        records = [{"date": "2026-08-06", "at": "09:00", "content": "정산 회신"}]
        html = weekly.build_html(records, dt.date(2026, 8, 3), dt.date(2026, 8, 7))
        self.assertIn("정산 회신", html)
        self.assertIn("08/06 (목)", html)
        self.assertIn("08/03 (월)", html)  # 빈 날도 행은 나온다

    def test_제목에_꺾쇠가_있어도_표가_깨지지_않는다(self):
        records = [{"date": "2026-08-06", "content": "<b>긴급</b> & 확인"}]
        html = weekly.build_html(records, dt.date(2026, 8, 3), dt.date(2026, 8, 7))
        self.assertIn("&lt;b&gt;긴급&lt;/b&gt;", html)
        self.assertNotIn("<b>긴급</b>", html)

    def test_합계가_맞는다(self):
        records = [
            {"date": "2026-08-04", "content": "A"},
            {"date": "2026-08-06", "content": "B"},
            {"date": "2026-08-06", "content": "C"},
        ]
        html = weekly.build_html(records, dt.date(2026, 8, 3), dt.date(2026, 8, 7))
        self.assertIn(">3</td>", html)

    def test_기간_밖_기록은_합계에_안_들어간다(self):
        records = [
            {"date": "2026-08-06", "content": "이번주"},
            {"date": "2026-07-30", "content": "지난주"},
        ]
        html = weekly.build_html(records, dt.date(2026, 8, 3), dt.date(2026, 8, 7))
        self.assertNotIn("지난주", html)
        self.assertIn(">1</td>", html)


if __name__ == "__main__":
    unittest.main()
