"""대기열 엑셀 생성과 되읽기.

왜 엑셀인가: 이미 매일 쓰는 도구고, **행 삭제가 가장 직관적인 "아니오"** 이기 때문이다.
새 프로그램을 배우게 하면 3주 뒤에 안 쓴다.
"""

from __future__ import annotations

import datetime as dt
import pathlib

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import QueueItem

FONT = "맑은 고딕"
ACCENT = "2563EB"
MUTED = "7B8794"
KEY_FILL = "F5F7FA"
MARK_FILL = "FFF8E1"

HEADERS = ["시각", "종류", "상대", "제목", "포함 사유", "구분", "업무 내용 / 메모", "_key"]
MARK_COLUMN = 6
CONTENT_COLUMN = 7
KEY_COLUMN = 8
FIRST_ROW = 4

# '아님'으로 받아들이는 표기. 사람마다 다르게 쓰기 때문에 넓게 잡는다.
# 빈칸은 업무로 본다 — 지우는 방식을 쓰던 사람이 그대로 써도 동작해야 한다.
NOT_WORK = {"아님", "x", "X", "ㅌ", "제외", "no", "N", "n"}


def is_not_work(mark) -> bool:
    return str(mark or "").strip() in NOT_WORK

_THIN = Side(style="thin", color="D2D6DC")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _title(ws, row: int, text: str, size: int = 14, color: str = "1F2933") -> None:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, size=size, bold=True, color=color)


def _header(ws, row: int) -> None:
    for index, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=row, column=index, value=name)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 24


def _write_rows(ws, start: int, items: list[QueueItem], *, with_hint: bool = False) -> None:
    for offset, item in enumerate(items):
        row = start + offset
        values = [
            item.at.strftime("%H:%M"),
            item.kind,
            item.counterpart,
            item.subject,
            item.hint if with_hint and item.hint else item.reason,
            "",
            "",
            item.key,
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center" if index in (1, 2, MARK_COLUMN) else "left",
                vertical="center",
            )
        ws.cell(row=row, column=MARK_COLUMN).fill = PatternFill("solid", fgColor=MARK_FILL)
        key_cell = ws.cell(row=row, column=KEY_COLUMN)
        key_cell.font = Font(name=FONT, size=8, color=MUTED)
        key_cell.fill = PatternFill("solid", fgColor=KEY_FILL)


def _mark_dropdown(ws, start: int, count: int) -> None:
    """'구분' 칸에 업무/아님 드롭다운을 건다.

    직접 타이핑해도 받아준다(NOT_WORK 참고). 드롭다운은 편의일 뿐,
    이걸로 입력을 강제하면 표기가 조금만 달라도 업무가 조용히 사라진다.
    """
    if count <= 0:
        return
    rule = DataValidation(type="list", formula1='"업무,아님"', allow_blank=True)
    ws.add_data_validation(rule)
    column = get_column_letter(MARK_COLUMN)
    rule.add(f"{column}{start}:{column}{start + count - 1}")


def write_queue(
    path: pathlib.Path,
    day: dt.date,
    queue: list[QueueItem],
    excluded: list[QueueItem] | None = None,
) -> pathlib.Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "오늘 업무"
    ws.sheet_view.showGridLines = False
    for column, width in {
        "A": 8, "B": 11, "C": 30, "D": 60, "E": 16, "F": 8, "G": 44, "H": 20
    }.items():
        ws.column_dimensions[column].width = width

    _title(ws, 1, f"{day.strftime('%Y-%m-%d (%a)')} 업무 대기열")
    note = ws.cell(
        row=2, column=1,
        value="업무가 아닌 줄은 '구분'을 [아님]으로 바꾸세요 (행을 지워도 같습니다). "
              "왜 아닌지를 '업무 내용 / 메모'에 적어두면 반복업무를 뽑을 때 근거로 씁니다.",
    )
    note.font = Font(name=FONT, size=9, color=MUTED)
    ws.merge_cells("A2:G2")

    _header(ws, 3)
    _write_rows(ws, FIRST_ROW, queue)
    _mark_dropdown(ws, FIRST_ROW, len(queue))
    ws.freeze_panes = f"A{FIRST_ROW}"
    ws.column_dimensions["H"].hidden = True  # _key는 손대지 말라는 뜻

    if excluded:
        # 조용히 사라지면 도구를 못 믿는다. 무엇을 뺐는지 보여준다.
        ws2 = wb.create_sheet("자동 제외됨")
        ws2.sheet_view.showGridLines = False
        for column, width in {
            "A": 8, "B": 11, "C": 30, "D": 60, "E": 26, "F": 8, "G": 44, "H": 20
        }.items():
            ws2.column_dimensions[column].width = width
        _title(ws2, 1, "자동으로 뺀 항목", size=12, color=MUTED)
        hint = ws2.cell(
            row=2, column=1,
            value="지금까지 항상 지우셨던 상대라 미리 뺐습니다. 그대로 두시면 계속 뺍니다. "
                  "잘못 뺀 게 있으면 이 행을 [오늘 업무] 시트로 옮겨주세요.",
        )
        hint.font = Font(name=FONT, size=9, color=MUTED)
        ws2.merge_cells("A2:G2")
        _header(ws2, 3)
        _write_rows(ws2, FIRST_ROW, excluded, with_hint=True)
        ws2.column_dimensions["H"].hidden = True

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_kept(path: pathlib.Path) -> dict[str, str]:
    """사용자가 업무로 남긴 행을 돌려준다. {key: 업무내용}

    '업무 내용'을 적었으면 그 값을, 안 적었으면 제목을 쓴다.

    [오늘 업무] 시트만 본다. [자동 제외됨]에 그대로 둔 행은 "빼는 데 동의"로 읽는다.
    예전에는 두 시트를 다 읽었는데, 그러면 자동 제외가 하루 만에 풀린다 —
    가만히 둔 항목이 "남긴 것"으로 집계돼 제외율이 1.0 밑으로 떨어지고,
    다음 날 그 뉴스레터가 대기열에 다시 올라오면서 업무 기록에도 들어갔다.
    되돌리려면 그 행을 [오늘 업무] 시트로 옮기면 된다.
    """
    wb = load_workbook(path, data_only=True)
    kept: dict[str, str] = {}

    if "오늘 업무" not in wb.sheetnames:
        return kept

    for row in wb["오늘 업무"].iter_rows(min_row=FIRST_ROW, values_only=True):
        if not row or len(row) < KEY_COLUMN:
            continue
        key = row[KEY_COLUMN - 1]
        if not key or not str(key).strip():
            continue
        if is_not_work(row[MARK_COLUMN - 1]):
            continue
        subject = str(row[3] or "").strip()
        override = str(row[CONTENT_COLUMN - 1] or "").strip()
        kept[str(key).strip()] = override or subject

    return kept


def read_memos(path: pathlib.Path) -> dict[str, tuple[bool, str]]:
    """모든 행의 판정과 메모를 돌려준다. {key: (업무인가, 메모)}

    read_kept 와 달리 '아님'으로 표시한 행의 메모도 가져온다.
    왜 업무가 아닌지가 반복업무를 추릴 때 쓸모가 있다 —
    "광고성 메일"과 "이번엔 조치 없는 공지"는 다르게 다뤄야 한다.
    """
    wb = load_workbook(path, data_only=True)
    memos: dict[str, tuple[bool, str]] = {}

    for sheet_name in ("오늘 업무", "자동 제외됨"):
        if sheet_name not in wb.sheetnames:
            continue
        is_queue = sheet_name == "오늘 업무"
        for row in wb[sheet_name].iter_rows(min_row=FIRST_ROW, values_only=True):
            if not row or len(row) < KEY_COLUMN:
                continue
            key = row[KEY_COLUMN - 1]
            if not key or not str(key).strip():
                continue
            is_work = is_queue and not is_not_work(row[MARK_COLUMN - 1])
            memo = str(row[CONTENT_COLUMN - 1] or "").strip()
            memos[str(key).strip()] = (is_work, memo)

    return memos
