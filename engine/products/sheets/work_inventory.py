#!/usr/bin/env python3
"""반복업무 인벤토리.

목적: **뭘 하고 있는지 먼저 보이게 한다. 자동화는 그 다음이다.**

정리가 안 된 상태에서 바로 자동화에 들어가면 가장 눈에 띄는 것부터 손대게 되고,
그건 대개 가장 값어치 없는 것이다. 목록을 만들면 어디에 시간이 새는지가 계산으로 나온다.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import Reference

from ..lib import (
    DATE, MONEY, MUTED, PERCENT, POSITIVE, QTY,
    add_sheet, autofill_rows, bar_chart, body_cell, color_scale, dropdown,
    guide_sheet, header_row, input_legend, kpi_card, new_workbook, note,
    section_title, set_widths, title_block,
)

FIRST, LAST = 6, 105        # 업무 100건
FREQ_FIRST, FREQ_LAST = 12, 17

RATE = "'설정'!$B$6"        # 시간당 가치
WEEKS = "'설정'!$B$7"       # 연간 근무주수
FREQ_TABLE = f"'설정'!$A${FREQ_FIRST}:$B${FREQ_LAST}"

W = "'업무목록'"
W_NAME = f"{W}!$A${FIRST}:$A${LAST}"
W_CATEGORY = f"{W}!$B${FIRST}:$B${LAST}"
W_HOURS = f"{W}!$H${FIRST}:$H${LAST}"
W_COST = f"{W}!$I${FIRST}:$I${LAST}"
W_SAVABLE = f"{W}!$L${FIRST}:$L${LAST}"
W_SCORE = f"{W}!$M${FIRST}:$M${LAST}"

CATEGORIES = ["문서·승인", "시스템 운영", "데이터 집계·보고", "커뮤니케이션", "기타"]
FREQUENCIES = ["매일", "주 2~3회", "주 1회", "격주", "월 1회", "분기 1회"]
AUTOMATABLE = ["완전 자동화", "반자동화", "불가"]


def _settings(wb: Workbook) -> None:
    ws = add_sheet(wb, "설정")
    set_widths(ws, {"A": 22, "B": 16, "C": 4, "D": 58})

    row = title_block(ws, 1, "설정", "두 값만 본인 기준으로 맞춰주세요.", span=4)
    row = input_legend(ws, row, span=4)
    header_row(ws, row, ["항목", "값", "", "설명"])
    row += 1
    assert row == 6, f"시간당 가치는 B6이어야 한다 (현재 B{row})"

    body_cell(ws, row, 1, "시간당 가치", align="left")
    body_cell(ws, row, 2, 30000, fmt=MONEY, is_input=True, bold=True)
    body_cell(ws, row, 4, "본인 시급 환산값. 절감 효과를 돈으로 보기 위한 기준일 뿐입니다",
              align="left", color=MUTED)
    row += 1
    body_cell(ws, row, 1, "연간 근무주수", align="left")
    body_cell(ws, row, 2, 48, fmt=QTY, is_input=True, bold=True)
    body_cell(ws, row, 4, "휴가를 뺀 실제 근무 주수. 연간 횟수 계산에 쓰입니다",
              align="left", color=MUTED)
    row += 1

    row += 2
    row = section_title(ws, row, "주기별 연간 횟수 (자동 계산)")
    header_row(ws, row, ["주기", "연간 횟수", "", ""])
    row += 1
    assert row == FREQ_FIRST, f"주기 테이블 첫 행이 {FREQ_FIRST}이어야 한다 (현재 {row})"

    for label, formula in [
        ("매일", f"=ROUND({WEEKS}*5,0)"),
        ("주 2~3회", f"=ROUND({WEEKS}*2.5,0)"),
        ("주 1회", f"={WEEKS}"),
        ("격주", f"=ROUND({WEEKS}/2,0)"),
        ("월 1회", "=12"),
        ("분기 1회", "=4"),
    ]:
        body_cell(ws, row, 1, label, align="left")
        body_cell(ws, row, 2, formula, fmt=QTY)
        row += 1

    note(ws, row + 2,
         "※ 근무주수를 바꾸면 모든 업무의 연간 소요시간이 자동으로 다시 계산됩니다.", span=4)


def _tasks(wb: Workbook) -> None:
    ws = add_sheet(wb, "업무목록")
    set_widths(ws, {
        "A": 34, "B": 16, "C": 12, "D": 12, "E": 20, "F": 22, "G": 11,
        "H": 13, "I": 14, "J": 11, "K": 13, "L": 14, "M": 12, "N": 28,
    })

    row = title_block(
        ws, 1, "업무 목록",
        "한 주 동안 실제로 한 반복 작업을 그때그때 한 줄씩 추가하세요. "
        "한 번에 다 채우려 하지 마세요 — 떠오르는 대로 적는 게 정확도보다 중요합니다.",
        span=14,
    )
    row = input_legend(ws, row, span=14)
    header_row(ws, row, [
        "업무명", "분류", "주기", "1회 소요\n(분)", "사용 시스템", "산출물",
        "연간 횟수", "연간 소요\n(시간)", "연간 비용", "자동화\n난이도",
        "자동화\n가능성", "절감 가능\n(시간)", "우선순위", "메모",
    ])
    assert row + 1 == FIRST

    # 자동화 가능성 → 절감 계수. 반자동은 절반만 줄어든다고 본다.
    savable = '=IF($K{row}="완전 자동화",1,IF($K{row}="반자동화",0.5,0))*$H{row}'

    autofill_rows(ws, FIRST, LAST, {
        1: {"is_input": True, "align": "left"},
        2: {"is_input": True, "align": "center"},
        3: {"is_input": True, "align": "center"},
        4: {"is_input": True, "fmt": QTY},
        5: {"is_input": True, "align": "left"},
        6: {"is_input": True, "align": "left"},
        7: {"formula": f'=IFERROR(VLOOKUP($C{{row}},{FREQ_TABLE},2,FALSE),0)', "fmt": QTY},
        8: {"formula": "=ROUND(N($D{row})*$G{row}/60,1)", "fmt": "#,##0.0"},
        9: {"formula": f"=ROUND($H{{row}}*{RATE},0)", "fmt": MONEY},
        10: {"is_input": True, "fmt": QTY, "align": "center"},
        11: {"is_input": True, "align": "center"},
        12: {"formula": savable, "fmt": "#,##0.0"},
        # 우선순위 = 절감 가능시간 ÷ 난이도. 적은 노력으로 많이 줄이는 것이 위로 온다.
        13: {"formula": "=IFERROR(ROUND($L{row}/MAX(1,N($J{row})),1),0)", "fmt": "#,##0.0"},
        14: {"is_input": True, "align": "left"},
    })

    dropdown(ws, f"B{FIRST}:B{LAST}", CATEGORIES)
    dropdown(ws, f"C{FIRST}:C{LAST}", FREQUENCIES)
    dropdown(ws, f"J{FIRST}:J{LAST}", ["1", "2", "3", "4", "5"])
    dropdown(ws, f"K{FIRST}:K{LAST}", AUTOMATABLE)
    color_scale(ws, f"H{FIRST}:H{LAST}")
    color_scale(ws, f"M{FIRST}:M{LAST}")
    ws.freeze_panes = f"B{FIRST}"

    ws.cell(row=FIRST, column=1, value="(예시) 주간 매출 집계 후 보고서 메일 발송")
    ws.cell(row=FIRST, column=2, value="데이터 집계·보고")
    ws.cell(row=FIRST, column=3, value="주 1회")
    ws.cell(row=FIRST, column=4, value=90)
    ws.cell(row=FIRST, column=10, value=2)
    ws.cell(row=FIRST, column=11, value="완전 자동화")
    ws.cell(row=FIRST, column=14, value="예시 행입니다. 지우고 쓰세요")

    note(ws, LAST + 2,
         "※ 자동화 난이도는 1(아주 쉬움) ~ 5(아주 어려움) 감으로 매기면 됩니다. "
         "정확할 필요 없습니다 — 순위를 가르는 용도입니다.\n"
         "※ 우선순위 = 절감 가능시간 ÷ 난이도. 적은 노력으로 많이 줄이는 것이 위로 옵니다.",
         span=14)


def _priority(wb: Workbook) -> None:
    ws = add_sheet(wb, "우선순위", tab_color="7C3AED")
    set_widths(ws, {
        "A": 34, "B": 16, "C": 13, "D": 14, "E": 13, "F": 12, "G": 4, "H": 15, "I": 15,
    })

    filled = f'COUNTIF({W_NAME},"<>")'

    row = title_block(ws, 1, "우선순위",
                      "업무목록에 입력하면 자동으로 채워집니다. 손댈 곳 없습니다.", span=9)

    for idx, (label, formula, fmt) in enumerate([
        ("등록 업무 수", f"={filled}", QTY),
        ("연간 총 소요", f"=ROUND(SUM({W_HOURS}),0)", "#,##0\"시간\""),
        ("연간 총 비용", f"=SUM({W_COST})", MONEY),
        ("절감 가능 시간", f"=ROUND(SUM({W_SAVABLE}),0)", "#,##0\"시간\""),
        ("절감 가능 금액", f"=ROUND(SUM({W_SAVABLE})*{RATE},0)", MONEY),
        ("절감 비율", f"=IFERROR(SUM({W_SAVABLE})/SUM({W_HOURS}),0)", PERCENT),
    ]):
        kpi_card(ws, row, idx + 1, label, formula, fmt)
    row += 4

    # ── 자동화 우선순위 상위 15 ─────────────────────────
    row = section_title(ws, row, "① 여기부터 자동화하세요 — 우선순위 상위 15")
    header_row(ws, row, [
        "업무명", "분류", "연간 소요", "절감 가능", "난이도", "우선순위",
    ])
    row += 1
    top_start = row
    for offset in range(15):
        r = top_start + offset
        rank = f"LARGE({W_SCORE},{offset + 1})"
        idx = f"MATCH({rank},{W_SCORE},0)"
        guard = f'IF(OR({filled}<{offset + 1},{rank}<=0),"",{{expr}})'
        body_cell(ws, r, 1, "=" + guard.format(expr=f"INDEX({W_NAME},{idx})"), align="left")
        body_cell(ws, r, 2, "=" + guard.format(expr=f"INDEX({W_CATEGORY},{idx})"), align="center")
        body_cell(ws, r, 3, "=" + guard.format(expr=f"INDEX({W_HOURS},{idx})"), fmt="#,##0.0")
        body_cell(ws, r, 4, "=" + guard.format(expr=f"INDEX({W_SAVABLE},{idx})"), fmt="#,##0.0")
        body_cell(ws, r, 5, "=" + guard.format(
            expr=f"INDEX({W}!$J${FIRST}:$J${LAST},{idx})"), fmt=QTY)
        body_cell(ws, r, 6, "=" + guard.format(expr=rank), fmt="#,##0.0", bold=True)
    top_end = top_start + 14
    color_scale(ws, f"F{top_start}:F{top_end}")

    bar_chart(ws, "우선순위 상위",
              Reference(ws, min_col=6, min_row=top_start - 1, max_row=top_end),
              Reference(ws, min_col=1, min_row=top_start, max_row=top_end),
              anchor=f"H{top_start - 1}", height=11)
    row = top_end + 3

    # ── 분류별 ──────────────────────────────────────────
    row = section_title(ws, row, "② 분류별 — 어디에 시간이 몰려 있는가")
    header_row(ws, row, ["분류", "업무 수", "연간 소요", "연간 비용", "절감 가능", "비중"])
    cat_head = row
    row += 1
    cat_start = row
    for offset, category in enumerate(CATEGORIES):
        r = cat_start + offset
        body_cell(ws, r, 1, category, align="left")
        body_cell(ws, r, 2, f"=COUNTIF({W_CATEGORY},$A{r})", fmt=QTY)
        body_cell(ws, r, 3, f"=ROUND(SUMIF({W_CATEGORY},$A{r},{W_HOURS}),1)", fmt="#,##0.0")
        body_cell(ws, r, 4, f"=SUMIF({W_CATEGORY},$A{r},{W_COST})", fmt=MONEY)
        body_cell(ws, r, 5, f"=ROUND(SUMIF({W_CATEGORY},$A{r},{W_SAVABLE}),1)", fmt="#,##0.0")
        body_cell(ws, r, 6, f"=IFERROR($C{r}/SUM({W_HOURS}),0)", fmt=PERCENT)
    cat_end = cat_start + len(CATEGORIES) - 1
    color_scale(ws, f"C{cat_start}:C{cat_end}")
    row = cat_end + 3

    # ── 자동화 가능성별 ─────────────────────────────────
    row = section_title(ws, row, "③ 자동화 가능성별 — 손댈 수 있는 게 얼마나 되는가")
    header_row(ws, row, ["가능성", "업무 수", "연간 소요", "절감 가능", "", ""])
    row += 1
    auto_start = row
    for offset, label in enumerate(AUTOMATABLE):
        r = auto_start + offset
        target = f"{W}!$K${FIRST}:$K${LAST}"
        body_cell(ws, r, 1, label, align="left")
        body_cell(ws, r, 2, f"=COUNTIF({target},$A{r})", fmt=QTY)
        body_cell(ws, r, 3, f"=ROUND(SUMIF({target},$A{r},{W_HOURS}),1)", fmt="#,##0.0")
        body_cell(ws, r, 4, f"=ROUND(SUMIF({target},$A{r},{W_SAVABLE}),1)", fmt="#,##0.0")
        row += 1

    note(ws, row + 1,
         "※ 상위 1~2개만 먼저 자동화하세요. 한꺼번에 다 하려 들면 아무것도 안 끝납니다.\n"
         "※ '불가'로 분류한 업무가 많다면, 정말 불가인지 다시 보세요. "
         "대개는 '방법을 모르는 것'이지 '불가능한 것'이 아닙니다.", span=9)


def build() -> Workbook:
    wb = new_workbook()
    _settings(wb)
    _tasks(wb)
    _priority(wb)
    guide_sheet(wb, "반복업무 인벤토리", [
        ("이 시트가 하는 일", [
            "내가 무슨 일에 시간을 쓰고 있는지를 눈에 보이게 만듭니다.",
            "그리고 어떤 것부터 자동화해야 효과가 큰지 계산해서 순위로 보여줍니다.",
        ]),
        ("왜 자동화보다 목록이 먼저인가", [
            "정리가 안 된 상태에서 바로 자동화에 들어가면 가장 눈에 띄는 것부터 손대게 됩니다.",
            "그런데 눈에 띄는 것과 시간을 많이 먹는 것은 대개 다릅니다.",
            "목록을 만들면 어디에 시간이 새는지가 계산으로 나옵니다.",
        ]),
        ("처음 한 번만", [
            "[설정]에서 시간당 가치와 연간 근무주수를 확인합니다.",
            "시간당 가치는 절감 효과를 돈으로 환산해 보기 위한 기준일 뿐입니다. 대충 넣어도 됩니다.",
        ]),
        ("한 주 동안", [
            "반복 작업을 할 때마다 [업무목록]에 한 줄씩 추가합니다.",
            "★ 한 번에 다 채우려 하지 마세요. 떠오르는 대로 적는 게 정확도보다 중요합니다.",
            "1회 소요시간은 감으로 적어도 됩니다. 순위를 가르는 용도입니다.",
            "자동화 난이도는 1(아주 쉬움) ~ 5(아주 어려움)로 매깁니다.",
        ]),
        ("한 주 뒤", [
            "[우선순위] 시트를 봅니다.",
            "① 상위 1~2개만 먼저 자동화하세요. 한꺼번에 다 하려 들면 아무것도 안 끝납니다.",
            "② 분류별로 보면 어느 영역에 시간이 몰려 있는지 보입니다.",
            "③ '불가'가 많다면 정말 불가인지 다시 보세요. 대개는 '방법을 모르는 것'입니다.",
        ]),
        ("계산 방식 (검증용)", [
            "연간 횟수 = 주기에 따라 [설정]의 표에서 자동 조회",
            "연간 소요시간 = 1회 소요(분) × 연간 횟수 ÷ 60",
            "연간 비용 = 연간 소요시간 × 시간당 가치",
            "절감 가능시간 = 연간 소요시간 × (완전 자동화 1.0 / 반자동화 0.5 / 불가 0)",
            "우선순위 = 절감 가능시간 ÷ 자동화 난이도",
            "  → 적은 노력으로 많이 줄이는 것이 위로 옵니다.",
        ]),
        ("⚠ 주의", [
            "회사 데이터·접속정보·사내 시스템 경로는 이 파일에 적지 마세요.",
            "업무명과 시스템 이름 정도면 충분합니다.",
        ]),
    ])
    return wb
