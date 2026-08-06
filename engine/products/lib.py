#!/usr/bin/env python3
"""엑셀 제품 빌더 공용 라이브러리.

제품을 하나씩 손으로 만들지 않는다. 주간 Routine이 이 라이브러리로 카탈로그를
계속 늘리기 때문에, 스타일과 구조는 여기 한 곳에서만 정의한다.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

# ── 디자인 토큰 ─────────────────────────────────────────────
FONT = "맑은 고딕"

INK = "1F2933"
MUTED = "7B8794"
ACCENT = "2563EB"
ACCENT_SOFT = "DBEAFE"
SURFACE = "F5F7FA"
LINE = "D2D6DC"
POSITIVE = "047857"
NEGATIVE = "B91C1C"
WARN_FILL = "FEF3C7"
INPUT_FILL = "FFFBEB"   # 노랑 — 사용자가 채우는 칸
REVIEW_FILL = "EFF6FF"  # 파랑 — 내가 채워주는 칸 (사용자가 판단할 수 없는 것)

MONEY = '#,##0"원"'
MONEY_PLAIN = "#,##0"
USD = '"$"#,##0.00'
PERCENT = "0.0%"
DATE = "yyyy-mm-dd"
QTY = "#,##0"

_THIN = Side(style="thin", color=LINE)
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def new_workbook() -> Workbook:
    """기본 시트가 제거된 빈 워크북."""
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def add_sheet(wb: Workbook, title: str, tab_color: str = ACCENT) -> Worksheet:
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    return ws


def set_widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def title_block(ws: Worksheet, row: int, title: str, subtitle: str = "", span: int = 8) -> int:
    """제목 + 설명. 다음에 쓸 행 번호를 돌려준다."""
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(name=FONT, size=16, bold=True, color=INK)
    ws.row_dimensions[row].height = 26
    row += 1
    if subtitle:
        cell = ws.cell(row=row, column=1, value=subtitle)
        cell.font = Font(name=FONT, size=10, color=MUTED)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30
        row += 1
    return row + 1


def section_title(ws: Worksheet, row: int, text: str, color: str = ACCENT) -> int:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, size=12, bold=True, color=color)
    return row + 1


def header_row(ws: Worksheet, row: int, headers: list[str], start_col: int = 1) -> None:
    fill = PatternFill("solid", fgColor=ACCENT)
    for offset, text in enumerate(headers):
        cell = ws.cell(row=row, column=start_col + offset, value=text)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def body_cell(
    ws: Worksheet,
    row: int,
    col: int,
    value=None,
    *,
    fmt: str | None = None,
    is_input: bool = False,
    is_review: bool = False,
    bold: bool = False,
    align: str = "right",
    color: str = INK,
):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(name=FONT, size=10, bold=bold, color=color)
    cell.border = BORDER
    cell.alignment = Alignment(horizontal=align, vertical="center")
    if fmt:
        cell.number_format = fmt
    if is_input:
        cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
    elif is_review:
        cell.fill = PatternFill("solid", fgColor=REVIEW_FILL)
    return cell


def note(ws: Worksheet, row: int, text: str, span: int = 8, color: str = MUTED) -> int:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, size=9, color=color, italic=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    return row + 1


def kpi_card(ws: Worksheet, row: int, col: int, label: str, formula: str, fmt: str = MONEY) -> None:
    """대시보드용 지표 카드 (라벨 위 / 값 아래)."""
    lab = ws.cell(row=row, column=col, value=label)
    lab.font = Font(name=FONT, size=9, bold=True, color=MUTED)
    lab.fill = PatternFill("solid", fgColor=SURFACE)
    lab.alignment = Alignment(horizontal="center", vertical="center")
    lab.border = BORDER

    val = ws.cell(row=row + 1, column=col, value=formula)
    val.font = Font(name=FONT, size=13, bold=True, color=ACCENT)
    val.fill = PatternFill("solid", fgColor=SURFACE)
    val.alignment = Alignment(horizontal="center", vertical="center")
    val.number_format = fmt
    val.border = BORDER
    ws.row_dimensions[row + 1].height = 24


def dropdown(
    ws: Worksheet,
    cell_range: str,
    options: list[str] | None = None,
    *,
    source: str | None = None,
) -> None:
    """목록 선택을 건다.

    `options`는 고정 목록, `source`는 다른 시트의 범위 참조다.
    범위 참조를 쓰면 사용자가 그 시트에 항목을 추가할 때 목록이 자동으로 따라온다.
    """
    if source:
        formula = source
    elif options:
        formula = '"' + ",".join(options) + '"'
    else:
        raise ValueError("options 또는 source 중 하나는 필요하다")

    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def color_negative(ws: Worksheet, cell_range: str) -> None:
    """음수를 빨갛게 — 손실 구간을 눈으로 잡기 위한 것."""
    ws.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            font=Font(name=FONT, size=10, bold=True, color=NEGATIVE),
            fill=PatternFill("solid", fgColor="FEE2E2"),
        ),
    )


def color_scale(ws: Worksheet, cell_range: str) -> None:
    ws.conditional_formatting.add(
        cell_range,
        ColorScaleRule(
            start_type="min", start_color="FEE2E2",
            mid_type="percentile", mid_value=50, mid_color="FFFFFF",
            end_type="max", end_color="D1FAE5",
        ),
    )


def bar_chart(ws: Worksheet, title: str, data_ref: Reference, cats_ref: Reference, anchor: str,
              width: float = 18, height: float = 8) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.title = title
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.width, chart.height = width, height
    chart.legend = None
    ws.add_chart(chart, anchor)


def pie_chart(ws: Worksheet, title: str, data_ref: Reference, cats_ref: Reference, anchor: str,
              width: float = 12, height: float = 8) -> None:
    chart = PieChart()
    chart.title = title
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    chart.width, chart.height = width, height
    ws.add_chart(chart, anchor)


def guide_sheet(wb: Workbook, product_name: str, sections: list[tuple[str, list[str]]]) -> Worksheet:
    """사용법 시트.

    환불 요청의 대부분은 '어떻게 쓰는지 모르겠다'에서 나온다. 그래서 모든 제품에 넣는다.
    """
    ws = add_sheet(wb, "📖 사용법", tab_color=MUTED)
    set_widths(ws, {"A": 4, "B": 96})
    ws.cell(row=1, column=2, value=product_name).font = Font(
        name=FONT, size=16, bold=True, color=INK
    )
    row = 3
    for heading, lines in sections:
        cell = ws.cell(row=row, column=2, value=heading)
        cell.font = Font(name=FONT, size=11, bold=True, color=ACCENT)
        cell.fill = PatternFill("solid", fgColor=ACCENT_SOFT)
        cell.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[row].height = 22
        row += 1
        for line in lines:
            cell = ws.cell(row=row, column=2, value=line)
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
            ws.row_dimensions[row].height = 30 if len(line) > 60 else 18
            row += 1
        row += 1
    return ws


def input_legend(ws: Worksheet, row: int, span: int = 8, *, with_review: bool = False) -> int:
    """색 범례. 어느 칸을 채워야 하는지 안 적으면 반드시 헷갈린다.

    `with_review=True`면 "내가 채워주는 칸"까지 3색으로 설명한다. 사용자가 판단할 수
    없는 항목(예: 자동화 난이도)을 억지로 물어보지 않기 위한 구분이다.
    """
    text = (
        "  노란색 = 직접 입력하세요   ·   파란색 = 제가 채워드립니다 (판단이 필요한 칸)"
        "   ·   나머지 = 자동 계산"
        if with_review
        else "  노란색 칸만 입력하세요. 나머지는 자동 계산됩니다."
    )
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = Font(name=FONT, size=9, bold=True, color="92400E")
    cell.fill = PatternFill("solid", fgColor=WARN_FILL)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 20
    return row + 1


def autofill_rows(ws: Worksheet, first_row: int, last_row: int, col_specs: dict[int, dict]) -> None:
    """빈 입력 행을 서식만 갖춘 채로 미리 깔아둔다.

    col_specs: {열번호: {"fmt": ..., "is_input": bool, "formula": "=...{row}..."}}
    formula의 {row}는 실제 행 번호로 치환된다.
    """
    for r in range(first_row, last_row + 1):
        for col, spec in col_specs.items():
            formula = spec.get("formula")
            value = formula.format(row=r) if formula else None
            body_cell(
                ws, r, col, value,
                fmt=spec.get("fmt"),
                is_input=spec.get("is_input", False),
                is_review=spec.get("is_review", False),
                align=spec.get("align", "right"),
            )


def col(idx: int) -> str:
    return get_column_letter(idx)
