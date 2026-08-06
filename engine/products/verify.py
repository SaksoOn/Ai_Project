#!/usr/bin/env python3
"""시트 수식 검증.

빌드된 xlsx에 테스트 데이터를 넣고 수식을 **실제로 계산**시킨 뒤,
손으로 계산한 기대값과 대조한다.

숫자가 틀린 시트는 잘못된 판단으로 이어진다. 특히 매매 복기는 그 잘못된 판단이
돈으로 이어진다. 눈으로 하는 검수로는 못 잡는다.

실행: python3 -m engine.products.verify [시트ID ...]
"""

from __future__ import annotations

import datetime as dt
import pathlib
import shutil
import sys
import tempfile

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[2]
DIST = ROOT / "dist"
TOLERANCE = 0.5  # 원 단위 반올림 오차만 허용. 1.0이면 건수 오류(0 vs 1)를 놓친다

EXCEL_ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")


class Sheet:
    """계산 결과에서 셀 값을 꺼내는 얇은 래퍼."""

    def __init__(self, solution: dict, filename: str, sheet: str):
        self._sol = solution
        self._prefix = f"'[{filename}]{sheet}'!"

    def __getitem__(self, cell: str):
        entry = self._sol.get(self._prefix + cell) or self._sol.get(self._prefix.upper() + cell)
        if entry is None:
            return None
        value = getattr(entry, "value", entry)
        try:
            return value[0][0]
        except (TypeError, IndexError, KeyError):
            return value


def calculate(path: pathlib.Path) -> dict:
    import formulas  # 무겁다. 검증할 때만 불러온다.

    return formulas.ExcelModel().loads(str(path)).finish().calculate()


def scan_errors(solution: dict, prefixes: tuple[str, ...], limit: int = 10) -> list[str]:
    """지정한 시트에서 엑셀 오류값을 훑는다.

    `prefixes`로 범위를 좁히는 이유: 계산 엔진이 EOMONTH 같은 일부 함수를 지원하지
    않아 거짓 양성이 나온다. 실제로 쓰는 핵심 시트만 본다.
    """
    found = []
    for key, entry in solution.items():
        if not any(p in key for p in prefixes):
            continue
        value = getattr(entry, "value", None)
        try:
            cell = value[0][0]
        except (TypeError, IndexError, KeyError):
            continue
        if isinstance(cell, str) and cell in EXCEL_ERRORS:
            found.append(f"{key.split(']')[-1]} = {cell}")
        if len(found) >= limit:
            break
    return found


class Checker:
    def __init__(self) -> None:
        self.failures: list[str] = []

    def __call__(self, label: str, actual, expected) -> None:
        if actual is None:
            print(f"    ✗ {label}: 셀이 계산되지 않음")
            self.failures.append(f"{label}: 계산되지 않음")
            return
        if isinstance(expected, str):
            ok, shown = str(actual).strip() == expected, actual
        else:
            try:
                shown = round(float(actual), 4)
                ok = abs(float(actual) - float(expected)) <= TOLERANCE
            except (TypeError, ValueError):
                shown, ok = actual, False
        print(f"    {'✓' if ok else '✗'} {label}: {shown} (기대 {expected})")
        if not ok:
            self.failures.append(f"{label}: {actual!r} ≠ {expected!r}")


def _stage(name: str, workdir: pathlib.Path) -> pathlib.Path | None:
    src = DIST / name
    if not src.exists():
        return None
    staged = workdir / src.name
    shutil.copy(src, staged)
    return staged


# ── 매매 복기 ────────────────────────────────────────────────

def verify_trade_review(workdir: pathlib.Path) -> list[str]:
    staged = _stage("매매복기.xlsx", workdir)
    if staged is None:
        return ["빌드 산출물 없음: dist/매매복기.xlsx"]

    wb = load_workbook(staged)
    watch, trades = wb["관심종목"], wb["매매기록"]
    watch["A6"], watch["A7"] = "종목A", "종목B"

    # 매매 1 — 이익, 31일 보유, 매수기준 지킴
    trades["A6"], trades["B6"], trades["C6"] = "종목A", dt.date(2026, 1, 10), dt.date(2026, 2, 10)
    trades["D6"], trades["E6"], trades["F6"] = 100, 10000, 12000
    trades["U6"], trades["V6"] = "예", "예"

    # 매매 2 — 손실, 120일 보유(더 오래!), 매수기준 어김
    trades["A7"], trades["B7"], trades["C7"] = "종목B", dt.date(2026, 1, 5), dt.date(2026, 5, 5)
    trades["D7"], trades["E7"], trades["F7"] = 50, 20000, 16000
    trades["U7"], trades["V7"] = "아니오", "아니오"
    wb.save(staged)

    print("  수식 계산 중…")
    solution = calculate(staged)
    t = Sheet(solution, staged.name, "매매기록")
    d = Sheet(solution, staged.name, "진단")

    # ── 손계산 ────────────────────────────────────
    fee_rate, tax_rate, bench = 0.00015, 0.0018, 0.07

    buy1, sell1 = 100 * 10000, 100 * 12000            # 1,000,000 / 1,200,000
    fee1 = round((buy1 + sell1) * fee_rate)           # 330
    tax1 = round(sell1 * tax_rate)                    # 2,160
    pl1 = sell1 - buy1 - fee1 - tax1                  # 197,510
    hold1 = 31
    bench1 = round(buy1 * hold1 * bench / 365)        # 5,945

    buy2, sell2 = 50 * 20000, 50 * 16000              # 1,000,000 / 800,000
    fee2 = round((buy2 + sell2) * fee_rate)           # 270
    tax2 = round(sell2 * tax_rate)                    # 1,440
    pl2 = sell2 - buy2 - fee2 - tax2                  # -201,710
    hold2 = 120
    bench2 = round(buy2 * hold2 * bench / 365)        # 23,014

    total_pl = pl1 + pl2                              # -4,200
    total_cost = fee1 + tax1 + fee2 + tax2            # 4,200 — 손실이 정확히 비용만큼이다

    check = Checker()
    print("  매매 1 (이익)")
    check("매수금액(G6)", t["G6"], buy1)
    check("매도금액(H6)", t["H6"], sell1)
    check("수수료(I6)", t["I6"], fee1)
    check("거래세(J6)", t["J6"], tax1)
    check("실현손익(K6)", t["K6"], pl1)
    check("보유일(M6)", t["M6"], hold1)
    check("결과(N6)", t["N6"], "이익")
    check("벤치마크(Q6)", t["Q6"], bench1)
    check("초과성과(R6)", t["R6"], pl1 - bench1)

    print("  매매 2 (손실, 더 오래 보유)")
    check("실현손익(K7)", t["K7"], pl2)
    check("보유일(M7)", t["M7"], hold2)
    check("결과(N7)", t["N7"], "손실")
    check("초과성과(R7)", t["R7"], pl2 - bench2)

    print("  진단 — KPI")
    check("총 실현손익(A5)", d["A5"], total_pl)
    check("수수료+세금(B5)", d["B5"], total_cost)
    check("승률(C5)", d["C5"], 0.5)
    check("손익비(D5)", d["D5"], round(pl1 / abs(pl2), 4))
    check("지수 대비(E5)", d["E5"], total_pl - (bench1 + bench2))
    check("종료 매매(F5)", d["F5"], 2)

    print("  진단 ① 처분효과 — 손실을 더 오래 들고 있는가")
    check("이익 평균 보유일(C10)", d["C10"], hold1)
    check("손실 평균 보유일(C11)", d["C11"], hold2)
    verdict = d["A12"]
    ok = isinstance(verdict, str) and "⚠" in verdict
    print(f"    {'✓' if ok else '✗'} 경고 문구: {str(verdict)[:60]}")
    if not ok:
        check.failures.append("처분효과 경고가 떠야 하는데 안 떴다")

    print("  진단 ② 규율 — 기준 지킴 vs 어김")
    check("매수기준 지킴 건수(B17)", d["B17"], 1)
    check("매수기준 지킴 손익(C17)", d["C17"], pl1)
    check("매수기준 어김 건수(B18)", d["B18"], 1)
    check("매수기준 어김 손익(C18)", d["C18"], pl2)

    print("  진단 ③ 비용")
    check("총 수수료(B26)", d["B26"], fee1 + fee2)
    check("총 거래세(B27)", d["B27"], tax1 + tax2)
    check("비용이 없었다면(B30)", d["B30"], total_pl + total_cost)

    print("  빈 행 안전성 (오류가 뜨면 모든 집계가 틀어진다)")
    for cell in ("G8", "K8", "N8", "R8"):
        value = t[cell]
        ok = value in (0, None, "") or (isinstance(value, (int, float)) and float(value) == 0)
        print(f"    {'✓' if ok else '✗'} 빈 행 {cell}: {value!r}")
        if not ok:
            check.failures.append(f"빈 행 {cell}에 오류값: {value!r}")

    errors = scan_errors(solution, ("매매기록", "관심종목"))
    check.failures.extend(f"오류 셀: {e}" for e in errors)
    for e in errors:
        print(f"    ✗ {e}")

    return check.failures


# ── 반복업무 인벤토리 ────────────────────────────────────────

def verify_work_inventory(workdir: pathlib.Path) -> list[str]:
    staged = _stage("반복업무_인벤토리.xlsx", workdir)
    if staged is None:
        return ["빌드 산출물 없음: dist/반복업무_인벤토리.xlsx"]

    wb = load_workbook(staged)
    tasks = wb["업무목록"]

    # 업무 1 — 주 1회 90분, 완전 자동화, 난이도 2 (예시 행을 덮어쓴다)
    tasks["A6"], tasks["B6"], tasks["C6"] = "주간 보고서", "데이터 집계·보고", "주 1회"
    tasks["D6"] = 90
    tasks["E6"] = tasks["F6"] = tasks["G6"] = tasks["P6"] = None
    tasks["K6"], tasks["L6"] = "완전 자동화", 2   # ← 파란 칸: 내가 채우는 자리

    # 업무 2 — 매일 15분, 반자동화, 난이도 3
    tasks["A7"], tasks["B7"], tasks["C7"] = "메일 분류", "커뮤니케이션", "매일"
    tasks["D7"] = 15
    tasks["K7"], tasks["L7"] = "반자동화", 3
    wb.save(staged)

    print("  수식 계산 중…")
    solution = calculate(staged)
    w = Sheet(solution, staged.name, "업무목록")
    p = Sheet(solution, staged.name, "우선순위")

    # ── 손계산 ────────────────────────────────────
    rate, weeks = 30000, 48

    freq1 = weeks                       # 주 1회 = 48
    hours1 = round(90 * freq1 / 60, 1)  # 72.0
    cost1 = round(hours1 * rate)        # 2,160,000
    save1 = hours1 * 1.0                # 72.0 (완전 자동화)
    score1 = round(save1 / 2, 1)        # 36.0

    freq2 = weeks * 5                   # 매일 = 240
    hours2 = round(15 * freq2 / 60, 1)  # 60.0
    cost2 = round(hours2 * rate)        # 1,800,000
    save2 = hours2 * 0.5                # 30.0 (반자동화)
    score2 = round(save2 / 3, 1)        # 10.0

    check = Checker()
    print("  업무 1 — 주 1회 90분, 완전 자동화")
    check("연간 횟수(H6)", w["H6"], freq1)
    check("연간 소요시간(I6)", w["I6"], hours1)
    check("연간 비용(J6)", w["J6"], cost1)
    check("절감 가능(N6)", w["N6"], save1)
    check("우선순위(O6)", w["O6"], score1)

    print("  업무 2 — 매일 15분, 반자동화(계수 0.5)")
    check("연간 횟수(H7)", w["H7"], freq2)
    check("연간 소요시간(I7)", w["I7"], hours2)
    check("절감 가능(N7)", w["N7"], save2)
    check("우선순위(O7)", w["O7"], score2)

    print("  우선순위 — KPI")
    check("연간 총 소요(B5)", p["B5"], round(hours1 + hours2))
    check("연간 총 비용(C5)", p["C5"], cost1 + cost2)
    check("절감 가능 시간(D5)", p["D5"], round(save1 + save2))
    check("절감 가능 금액(E5)", p["E5"], round((save1 + save2) * rate))
    check("절감 비율(F5)", p["F5"], round((save1 + save2) / (hours1 + hours2), 4))

    print("  우선순위 — 상위 순위가 실제로 정렬되는가")
    check("1위 업무명(A10)", p["A10"], "주간 보고서")
    check("1위 점수(F10)", p["F10"], score1)
    check("2위 업무명(A11)", p["A11"], "메일 분류")
    check("2위 점수(F11)", p["F11"], score2)

    print("  빈 행 안전성")
    for cell in ("H8", "I8", "N8", "O8"):
        value = w[cell]
        ok = value in (0, None, "") or (isinstance(value, (int, float)) and float(value) == 0)
        print(f"    {'✓' if ok else '✗'} 빈 행 {cell}: {value!r}")
        if not ok:
            check.failures.append(f"빈 행 {cell}에 오류값: {value!r}")

    errors = scan_errors(solution, ("업무목록",))
    check.failures.extend(f"오류 셀: {e}" for e in errors)
    for e in errors:
        print(f"    ✗ {e}")

    return check.failures


VERIFIERS = {
    "trade-review": verify_trade_review,
    "work-inventory": verify_work_inventory,
}


def main(argv: list[str]) -> int:
    selected = {k: v for k, v in VERIFIERS.items() if not argv or k in argv}
    if not selected:
        print(f"검증기 없음: {', '.join(argv)}", file=sys.stderr)
        return 1

    all_failures: dict[str, list[str]] = {}
    with tempfile.TemporaryDirectory() as tmp:
        for sheet_id, verifier in selected.items():
            print(f"\n▶ {sheet_id}")
            sub = pathlib.Path(tmp) / sheet_id
            sub.mkdir(parents=True, exist_ok=True)
            try:
                failures = verifier(sub)
            except Exception as exc:  # noqa: BLE001
                import traceback
                traceback.print_exc(limit=4)
                failures = [f"검증 중 예외: {exc}"]
            if failures:
                all_failures[sheet_id] = failures

    print("\n" + "─" * 62)
    if all_failures:
        print("검증 실패")
        for sheet_id, failures in all_failures.items():
            print(f"\n{sheet_id}:")
            for f in failures:
                print(f"  - {f}")
        return 1
    print(f"검증 통과 — 시트 {len(selected)}종의 수식이 실제 계산 결과와 일치")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
